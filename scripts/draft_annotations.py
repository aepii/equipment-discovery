"""Draft, check and export the instrument annotations (the knowledge layer).

Usage (from the project folder):
  python scripts/draft_annotations.py check
      check data/vocabulary.json, data/techniques.json and data/annotations.json against each other
      and against data/dataset.xlsx
  python scripts/draft_annotations.py draft [--ids EQ05 EQ06] [--out data/annotations_draft.json]
      ask the LLM (LLM_PROVIDER in .env) to draft annotations, one call per instrument, using the
      controlled vocabulary and technique list; checks the draft and writes it to --out. It never
      overwrites data/annotations.json: copy reviewed entries over by hand.
  python scripts/draft_annotations.py export [--out data/instrument_capabilities.xlsx]
      write the annotations as a spreadsheet for review by the faculty mentor

annotations.json is the source of truth; the spreadsheet is generated from it. Corrections from the
review are made in the JSON files and committed as data changes.
"""

import argparse
import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

import core  # noqa: E402
import knowledge  # noqa: E402
import llm  # noqa: E402

DRAFT_PATH = ROOT / "data" / "annotations_draft.json"
EXPORT_PATH = ROOT / "data" / "instrument_capabilities.xlsx"


# --- draft ---

DRAFT_SYSTEM = """You draft notes about one research instrument for an equipment-discovery tool.
Describe what this TYPE of instrument typically does. The dataset does not list installed accessories,
fixtures or configurations: never say an accessory is installed unless the dataset notes say so.

Rules:
1. Use ONLY technique IDs from the technique list. If the instrument performs a technique that is not
   on the list, put it in "proposed_techniques" instead (a human will add it).
2. For each technique give "source": "dataset" if the dataset row (name or notes) states it, else
   "llm_draft"; and "confidence": "high" (standard for this instrument type), "medium" (common but
   depends on configuration) or "low" (only with optional hardware the dataset does not list).
3. Instruments that make or prepare samples (printers, mills, sectioning) measure nothing; say so in
   "limitations".
4. Accessories: what a measurement typically needs. "for" lists the technique IDs it is needed for.
   "status" is "listed_in_dataset" (with "source": "dataset") only if the dataset notes list it,
   else "to_confirm" (with "source": "llm_draft").

Return ONLY a JSON object with this shape:
{
  "id": "EQ19",
  "instrument": "the instrument name exactly as in the dataset",
  "techniques": [{"technique": "rotational_rheometry", "source": "dataset", "confidence": "high"}],
  "typical_samples": ["..."],
  "limitations": ["..."],
  "accessories": [{"item": "...", "for": ["rotational_rheometry"], "status": "to_confirm", "source": "llm_draft"}],
  "proposed_techniques": [{"name": "...", "measures": "..."}]
}"""


def technique_list() -> str:
    lines = []
    for tid, t in knowledge.techniques().items():
        measures = ", ".join(knowledge.property_label(p) for p in t["measures"]) or "nothing"
        lines.append(f"- {tid}: {t['name']} [{t['kind']}] measures: {measures}")
    return "\n".join(lines)


def draft_prompt(item: dict) -> str:
    row = (f"{item['id']} | {item['instrument']} | {item['manufacturer']} | {item['model']} | "
           f"notes: {item['notes'] or '(none)'}")
    return f"Technique list:\n{technique_list()}\n\nDataset row:\n{row}"


def draft_one(item: dict) -> dict:
    """One instrument's draft annotation, marked as an unreviewed draft."""
    parsed = core.parse_json(llm.complete(DRAFT_SYSTEM, draft_prompt(item), json_mode=True))
    parsed["id"], parsed["instrument"] = item["id"], item["instrument"]
    parsed.setdefault("techniques", [])
    parsed.setdefault("typical_samples", [])
    parsed.setdefault("limitations", [])
    parsed.setdefault("accessories", [])
    parsed["review"] = {"status": "draft", "reviewed_by": None, "reviewed_on": None, "comments": ""}
    return parsed


def draft(ids: list[str] | None, out: Path) -> list[str]:
    """Draft annotations, write them to out, and return the problems found in the draft."""
    equipment = [e for e in core.load_equipment() if not ids or e["id"] in ids]
    drafts, problems = [], []
    for item in equipment:
        try:
            drafts.append(draft_one(item))
            print(f"{item['id']} drafted")
        except Exception as exc:  # keep going; report at the end
            problems.append(f"{item['id']}: could not draft ({type(exc).__name__}: {exc})")
    problems += [p for p in knowledge.validate(equipment, annots=drafts) if "has no annotation" not in p]
    for d in drafts:
        for proposed in d.get("proposed_techniques", []):
            problems.append(f"{d['id']}: proposes a new technique {proposed} (add it to techniques.json if right)")
    out.write_text(json.dumps({"_about": "LLM draft; review before copying into annotations.json.",
                               "instruments": drafts}, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(f"Saved {out}")
    return problems


# --- export for review ---

GREY, YELLOW, HEADER = "D9D9D9", "FFF2CC", "BDD7EE"
DATASET_COLUMNS = ["ID", "Instrument", "Manufacturer", "Model", "Notes", "Location", "Contact"]
NOTE_COLUMNS = ["Techniques", "What it measures", "Destructive?", "Must come before", "Leads to",
                "Typical samples", "Limitations", "Accessories to confirm", "Installed per the dataset",
                "Source and confidence", "Review status", "Reviewer comments"]


def _join(items) -> str:
    return "; ".join(dict.fromkeys(i for i in items if i)) or "none"


def instrument_row(e: dict) -> list[str]:
    """One spreadsheet row: the dataset columns, then the notes, flattened to text."""
    a = knowledge.annotations()[e["id"]]
    techs = knowledge.techniques()
    own = knowledge.techniques_of(e["id"])
    measures = [knowledge.property_label(p) for t in own for p in techs[t]["measures"]]
    measures += [f"{knowledge.property_label(p)} (partly: {why})" for t in own
                 for p, why in techs[t].get("partial", {}).items()]
    destructive = [f"{techs[t]['name']}: {techs[t]['destructive'].replace('_', ' ')}" for t in own]
    before = [f"{knowledge.technique_name(s['technique'])} ({s['when']})" for t in own
              for s in techs[t]["requires_before"]]
    after = [knowledge.technique_name(s) for t in own for s in techs[t]["leads_to"]]
    confirm = [x["item"] for x in a["accessories"] if x["status"] == "to_confirm"]
    listed = [x["item"] for x in a["accessories"] if x["status"] == "listed_in_dataset"]
    sources = [f"{techs[t]['name']}: {own[t]['source']}, {own[t]['confidence']}" for t in own]
    return ([e["id"], e["instrument"], e["manufacturer"], e["model"], e["notes"], e["location"], e["contact"]]
            + [_join(techs[t]["name"] for t in own), _join(measures) if measures else "nothing (makes or prepares samples)",
               _join(destructive), _join(before), _join(after), _join(a["typical_samples"]),
               _join(a["limitations"]), _join(confirm), _join(listed), _join(sources),
               a["review"]["status"], a["review"].get("comments", "")])


def technique_row(tid: str, t: dict) -> list[str]:
    return [tid, t["name"], t["kind"], _join(knowledge.property_label(p) for p in t["measures"]),
            _join(f"{knowledge.property_label(p)}: {why}" for p, why in t.get("partial", {}).items()),
            t["destructive"].replace("_", " ") + (f" ({t['destructive_note']})" if t.get("destructive_note") else ""),
            _join(f"{knowledge.technique_name(s['technique'])} ({s['when']})" for s in t["requires_before"]),
            _join(knowledge.technique_name(s) for s in t["leads_to"]),
            _join(knowledge.instruments_for(tid)) if knowledge.instruments_for(tid) else "NOT AVAILABLE in the dataset",
            ""]


HOW_TO_REVIEW = [
    "Instrument capabilities: review copy",
    "Generated from data/annotations.json and data/techniques.json by scripts/draft_annotations.py export.",
    "Grey columns are Stefan's original dataset. The other columns are draft notes about what this TYPE of "
    "instrument typically does; nothing here says an accessory is installed unless 'Installed per the dataset' says so.",
    "Please write corrections in the 'Reviewer comments' column (Instruments) or the 'Comments' column (Techniques). "
    "The team updates the JSON files from them; edits elsewhere in this file are overwritten on the next export.",
    "Source: dataset = stated in the dataset; llm_draft = drafted with an LLM, unreviewed.",
    "Confidence: high = standard for this instrument type; medium = depends on configuration; "
    "low = only with optional hardware the dataset does not list.",
    "Techniques sheet: techniques marked NOT AVAILABLE are ones no listed instrument performs (e.g. micro-CT); "
    "the app uses them to say what is missing.",
]


def _style_sheet(sheet, fills: list[str], widths: list[int]) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    for cell, fill in zip(sheet[1], fills):
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor=fill)
    for row in sheet.iter_rows(min_row=2):
        for cell, fill in zip(row, fills):
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if fill == GREY:
                cell.fill = PatternFill("solid", fgColor=GREY)
    for i, width in enumerate(widths, start=1):
        sheet.column_dimensions[sheet.cell(1, i).column_letter].width = width
    sheet.freeze_panes = "C2"


def export(out: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    guide = wb.active
    guide.title = "How to review"
    for line in HOW_TO_REVIEW:
        guide.append([line])
    guide["A1"].font = Font(bold=True, size=14)
    guide.column_dimensions["A"].width = 140

    sheet = wb.create_sheet("Instruments")
    sheet.append(DATASET_COLUMNS + NOTE_COLUMNS)
    for e in core.load_equipment():
        sheet.append(instrument_row(e))
    _style_sheet(sheet, [GREY] * len(DATASET_COLUMNS) + [YELLOW] * len(NOTE_COLUMNS),
                 [7, 28, 16, 22, 30, 28, 20] + [32, 40, 32, 32, 28, 30, 40, 45, 30, 40, 12, 30])

    techs = wb.create_sheet("Techniques")
    techs.append(["ID", "Technique", "Kind", "Measures", "Partly", "Destructive?", "Must come before",
                  "Leads to", "Instruments", "Comments"])
    for tid, t in knowledge.techniques().items():
        techs.append(technique_row(tid, t))
    _style_sheet(techs, [HEADER] * 9 + [YELLOW], [30, 40, 16, 45, 45, 40, 40, 32, 24, 30])
    wb.save(out)
    print(f"Saved {out}")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = parser.add_subparsers(dest="command", required=True)
    sub.add_parser("check", help="check the knowledge files")
    p_draft = sub.add_parser("draft", help="draft annotations with the LLM")
    p_draft.add_argument("--ids", nargs="+", metavar="ID", help="only these instruments")
    p_draft.add_argument("--out", type=Path, default=DRAFT_PATH)
    p_export = sub.add_parser("export", help="write the review spreadsheet")
    p_export.add_argument("--out", type=Path, default=EXPORT_PATH)
    args = parser.parse_args()

    if args.command == "check":
        problems = knowledge.validate(list(core.load_equipment()))
        print("\n".join(problems) or f"OK: {len(knowledge.annotations())} instruments, "
                                      f"{len(knowledge.techniques())} techniques, "
                                      f"{len(knowledge.vocabulary()['properties'])} properties.")
        sys.exit(1 if problems else 0)
    if args.command == "draft":
        problems = draft(args.ids, args.out)
        print("\n".join(problems) or "The draft passes the checks.")
    if args.command == "export":
        export(args.out)


if __name__ == "__main__":
    main()
