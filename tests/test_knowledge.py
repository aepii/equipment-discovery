"""The knowledge layer: the data files, their checks, annotated mode and the drafting script."""

import copy
import importlib.util
import json
from pathlib import Path

import openpyxl
import pytest

import core
import knowledge
import llm
import run_eval

EQUIPMENT = list(core.load_equipment())


def load_script():
    path = Path(__file__).parent.parent / "scripts" / "draft_annotations.py"
    spec = importlib.util.spec_from_file_location("draft_annotations", path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


# --- the data files ---

def test_knowledge_files_pass_the_checks():
    assert knowledge.validate(EQUIPMENT) == []


def test_every_instrument_is_annotated_and_still_a_draft():
    assert list(knowledge.annotations()) == [e["id"] for e in EQUIPMENT]
    assert {knowledge.review_status(e["id"]) for e in EQUIPMENT} == {"draft"}


def test_micro_ct_and_gpc_are_not_available():
    unavailable = knowledge.unavailable_techniques()
    for tid in ("micro_computed_tomography", "gel_permeation_chromatography", "capillary_rheometry"):
        assert tid in unavailable
    assert "rotational_rheometry" not in unavailable


@pytest.mark.parametrize("eq_id", ["EQ01", "EQ02", "EQ03", "EQ04", "EQ05", "EQ16"])
def test_instruments_that_make_or_prepare_samples_measure_nothing(eq_id):
    assert all(not knowledge.techniques()[t]["measures"] for t in knowledge.techniques_of(eq_id))


def test_electron_microscopy_and_ultramicrotomy_are_destructive():
    techs = knowledge.techniques()
    assert techs["transmission_electron_microscopy"]["destructive"] == "yes"
    assert techs["ultramicrotomy"]["destructive"] == "yes"
    assert techs["micro_computed_tomography"]["destructive"] == "no"


def test_instron_compression_needs_load_cell_and_platens_to_confirm():
    accessories = " ".join(knowledge.accessories_for("EQ15", "compression_testing")).lower()
    assert "load cell" in accessories and "platens" in accessories


def test_only_dataset_facts_are_listed_as_installed():
    listed = [(a["id"], x["item"]) for a in knowledge.annotations().values()
              for x in a["accessories"] if x["status"] != "to_confirm"]
    assert listed == [("EQ01", "two low-temperature heads, one high-temperature head, one UV-light head")]


def test_ultramicrotomy_leads_to_tem_which_requires_it():
    techs = knowledge.techniques()
    assert "transmission_electron_microscopy" in techs["ultramicrotomy"]["leads_to"]
    assert "ultramicrotomy" in [s["technique"] for s in techs["transmission_electron_microscopy"]["requires_before"]]


# --- the checks catch mistakes ---

def annotations_with(change) -> list[dict]:
    annots = copy.deepcopy(list(knowledge.annotations().values()))
    change({a["id"]: a for a in annots})
    return annots


@pytest.mark.parametrize("change, problem", [
    (lambda a: a["EQ19"]["techniques"].append({"technique": "magic", "source": "llm_draft", "confidence": "high"}),
     "'magic' is not in techniques.json"),
    (lambda a: a["EQ19"].update(instrument="Rheometer"), "does not match the dataset"),
    (lambda a: a["EQ15"]["accessories"][0].update(status="listed_in_dataset"), "its source is not the dataset"),
    (lambda a: a["EQ15"]["accessories"][0].update(status="installed"), "accessory status 'installed'"),
    (lambda a: a["EQ15"]["accessories"][0].update({"for": ["rotational_rheometry"]}), "does not perform"),
    (lambda a: a["EQ15"]["techniques"][0].update(confidence="certain"), "confidence 'certain'"),
])
def test_validate_finds_problems_in_annotations(change, problem):
    problems = knowledge.validate(EQUIPMENT, annots=annotations_with(change))
    assert any(problem in p for p in problems), problems


def test_validate_finds_missing_annotations_and_unknown_properties():
    annots = [a for a in knowledge.annotations().values() if a["id"] != "EQ20"]
    assert "EQ20 has no annotation" in knowledge.validate(EQUIPMENT, annots=annots)
    techs = copy.deepcopy(knowledge.techniques())
    techs["dynamic_light_scattering"]["measures"].append("colour")
    assert any("'colour' is not in the vocabulary" in p for p in knowledge.validate(EQUIPMENT, techs=techs))


# --- annotated mode ---

def test_annotated_prompt_has_notes_and_unavailable_techniques_but_no_contacts():
    prompt = core.build_user_prompt("q", "annotated")
    assert "compression_testing" in prompt and "load cell" in prompt
    assert "micro_computed_tomography" in prompt.split("Techniques not available")[1]
    assert core.equipment_by_id()["EQ15"]["contact"] not in prompt
    assert "technique_id" in core.ANNOTATED_SYSTEM
    assert core.prompt_version("annotated") != core.prompt_version("structured")


def annotated(recs, **extra) -> dict:
    return core.read_answer("annotated", json.dumps({"recommendations": recs, "summary": "s", **extra}))


def test_annotated_answer_is_checked_against_the_notes():
    result = annotated([
        {"instrument_id": "EQ15", "technique_id": "compression_testing", "technique": "Compression",
         "why": "", "accessories_to_confirm": []},
        {"instrument_id": "EQ01", "technique_id": "compression_testing", "technique": "Compression",
         "why": "", "accessories_to_confirm": []},
        {"instrument_id": "EQ10", "technique_id": "energy_dispersive_xray_spectroscopy", "technique": "EDS",
         "why": "", "accessories_to_confirm": ["EDS detector"]},
        {"instrument_id": "EQ12", "technique": "Brightfield", "why": "", "accessories_to_confirm": []},
        {"instrument_id": "EQ99", "technique_id": "x", "technique": "?", "why": ""},
    ])
    instron, printer, esem, microscope = result["recommendations"]
    assert result["removed"] == ["EQ99"]
    assert instron["cautions"] == [] and instron["destructive"].startswith("depends")
    assert any("load cell" in a for a in instron["accessories_to_confirm"])  # filled in from the notes
    assert "not among this instrument's typical techniques" in printer["cautions"][0]
    assert "optional hardware" in esem["cautions"][0]
    assert esem["accessories_to_confirm"] == ["EDS detector"]  # the model's list is kept
    assert "No technique ID" in microscope["cautions"][0]
    assert instron["notes_status"] == "draft"


def test_not_available_that_is_available_is_flagged():
    result = annotated([], not_available=[
        {"technique_id": "micro_computed_tomography", "technique": "Micro-CT"},
        {"technique_id": "rotational_rheometry", "technique": "Rheology"},
        {"technique_id": None, "technique": "Something else"}])
    assert [n["available_on"] for n in result["not_available"]] == [[], ["EQ19"], []]


def test_structured_mode_is_not_checked_against_the_notes():
    rec = {"instrument_id": "EQ01", "technique_id": "compression_testing", "technique": "x", "why": ""}
    result = core.read_answer("structured", json.dumps({"recommendations": [rec], "summary": "s"}))
    assert "cautions" not in result["recommendations"][0]


def test_run_annotated_with_mock_model(monkeypatch):
    prompts = []
    real = llm.complete_with_meta
    monkeypatch.setattr(core, "complete_with_meta", lambda s, u, **kw: prompts.append(s) or real(s, u, **kw))
    result = core.run_annotated("I want to determine the compressive stiffness of a porous polymer scaffold.")
    assert result["mode"] == "annotated" and not result["error"]
    assert result["instrument_ids"] == ["EQ15"] and result["removed"] == ["EQ99"]
    assert prompts == [core.ANNOTATED_SYSTEM]


def test_eval_scores_annotated_accessories_per_instrument():
    question = {"expected_ids": ["EQ15"], "expected_not_available": ["micro.?CT"],
                "expected_accessories": {"EQ15": ["load cell", "platen"]}}
    result = annotated([{"instrument_id": "EQ15", "technique_id": "compression_testing", "technique": "c",
                         "why": "", "accessories_to_confirm": []}],
                       not_available=[{"technique_id": "micro_computed_tomography", "technique": "Micro-CT"}])
    s = run_eval.score(question, result)
    assert (s["expected_hits"], s["accessory_hits"], s["na_hits"], s["contacts_total"]) == (1, 2, 1, None)


# --- drafting script ---

def test_draft_writes_a_checked_draft_and_reports_proposals(monkeypatch, tmp_path):
    script = load_script()
    answer = {"id": "EQ19", "instrument": "wrong name", "techniques": [
        {"technique": "rotational_rheometry", "source": "dataset", "confidence": "high"},
        {"technique": "made_up", "source": "llm_draft", "confidence": "low"}],
        "typical_samples": ["gels"], "limitations": [], "accessories": [],
        "proposed_techniques": [{"name": "Extensional rheometry", "measures": "extensional viscosity"}]}
    monkeypatch.setattr(script.llm, "complete", lambda system, user, json_mode=False: json.dumps(answer))
    before = knowledge.ANNOTATIONS_PATH.read_bytes()
    out = tmp_path / "draft.json"
    problems = script.draft(["EQ19"], out)
    saved = json.loads(out.read_text(encoding="utf-8"))["instruments"]
    assert saved[0]["instrument"] == "Hybrid rheometer" and saved[0]["review"]["status"] == "draft"
    assert any("'made_up' is not in techniques.json" in p for p in problems)
    assert any("Extensional rheometry" in p for p in problems)
    assert not any("has no annotation" in p for p in problems)
    assert knowledge.ANNOTATIONS_PATH.read_bytes() == before


def test_draft_prompt_lists_techniques_and_the_dataset_row():
    script = load_script()
    prompt = script.draft_prompt(core.equipment_by_id()["EQ01"])
    assert "micro_computed_tomography" in prompt and "Two low-temperature heads" in prompt


def test_export_writes_one_row_per_instrument(tmp_path):
    script = load_script()
    out = tmp_path / "review.xlsx"
    script.export(out)
    wb = openpyxl.load_workbook(out)
    rows = list(wb["Instruments"].iter_rows(min_row=2, values_only=True))
    assert [r[0] for r in rows] == [e["id"] for e in EQUIPMENT]
    printer = dict(zip([c.value for c in wb["Instruments"][1]], rows[0]))
    assert printer["What it measures"] == "nothing (makes or prepares samples)"
    techniques = {r[0]: r for r in wb["Techniques"].iter_rows(min_row=2, values_only=True)}
    assert techniques["micro_computed_tomography"][8] == "NOT AVAILABLE in the dataset"
